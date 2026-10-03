#!/usr/bin/env python3
"""Build the branded PowerPoint and Word templates in the brand skill from assets/tokens.css.

    uv run --with python-pptx --with python-docx scripts/build_brand_templates.py
    python3 scripts/build_brand_skill.py      # then refresh the zip

Writes brand-skill/appleseed-brand/templates/appleseed-slides.pptx and
appleseed-document.docx. Both carry an Office theme built from the tokens —
Charcoal/Slate text colours, the categorical chart palette as Accent 1–6 (so a
chart inserted in PowerPoint or Word is on-brand by default), Manrope headings
and Poppins body as the theme fonts — plus styled layouts/styles and a few
sample slides/paragraphs showing each one. Office substitutes a fallback when
Manrope/Poppins aren't installed; the skill says to use Arial then.
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_design_tokens import TOKENS, parse_tokens, hexof  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
SKILL = REPO / "brand-skill" / "appleseed-brand"
OUT = SKILL / "templates"
LOGO, LOGO_WHITE = SKILL / "assets" / "logo.png", SKILL / "assets" / "logo-white.png"
TOK = {k: hexof(v)[1:] for k, v in parse_tokens(TOKENS.read_text()).items() if v[3] == 1}
HEAD, BODY = "Manrope", "Poppins"
ORG = "Hawaiʻi Appleseed Center for Law & Economic Justice"


def themed(xml: bytes) -> bytes:
    """Point an Office theme's colour and font schemes at the brand tokens."""
    s = xml.decode("utf-8")
    slots = {"dk1": "--ha-charcoal", "lt1": "--ha-white", "dk2": "--ha-slate", "lt2": "--ha-bg",
             **{f"accent{i}": f"--ha-chart-cat-{i}" for i in range(1, 7)},
             "hlink": "--ha-teal-deep", "folHlink": "--ha-slate"}
    for slot, tok in slots.items():
        s = re.sub(rf"(<a:{slot}>)(.*?)(</a:{slot}>)",
                   lambda m: f'{m.group(1)}<a:srgbClr val="{TOK[tok]}"/>{m.group(3)}', s, count=1, flags=re.S)
    s = re.sub(r'(<a:clrScheme name=")[^"]*', r"\1Hawaii Appleseed", s, count=1)
    s = re.sub(r'(<a:fontScheme name=")[^"]*', r"\1Hawaii Appleseed", s, count=1)
    # pitchFamily 34 = proportional sans: where the font is missing, Office substitutes a sans, not Times.
    s = re.sub(r'(<a:majorFont>\s*)<a:latin [^>]*/>', rf'\g<1><a:latin typeface="{HEAD}" pitchFamily="34" charset="0"/>', s, count=1)
    s = re.sub(r'(<a:minorFont>\s*)<a:latin [^>]*/>', rf'\g<1><a:latin typeface="{BODY}" pitchFamily="34" charset="0"/>', s, count=1)
    return s.encode("utf-8")


def retheme(part, rel_type):
    theme = part.part_related_by(rel_type)
    theme._blob = themed(theme.blob)


# ---------------- PowerPoint ----------------

def build_pptx(path: Path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.util import Emu, Inches, Pt

    rgb = lambda t: RGBColor.from_string(TOK[t])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    master = prs.slide_master
    retheme(master.part, RT.THEME)

    # Master text styles: headings Manrope bold Charcoal, body Poppins Charcoal.
    tx = master._element.find(qn("p:txStyles"))
    for style, bold in (("p:titleStyle", True), ("p:bodyStyle", False)):
        for d in tx.find(qn(style)).iter(qn("a:defRPr")):
            d.set("b", "1" if bold else "0")
    title_lvl1 = tx.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    title_lvl1.set("algn", "l")  # brand: left-aligned headlines
    title_lvl1.find(qn("a:defRPr")).set("sz", "3600")
    for shapes in [master.placeholders] + [l.placeholders for l in master.slide_layouts]:
        for ph in shapes:  # python-pptx's default template is 4:3; stretch placeholders to 16:9
            if ph.left is not None:  # set all four, or python-pptx zeroes the ones left unset
                l, t, w, h = ph.left, ph.top, ph.width, ph.height
                ph.left, ph.top, ph.width, ph.height = int(l * W / Inches(10)), t, int(w * W / Inches(10)), h

    def text_colour(ph, tok):
        """Set a layout placeholder's inherited text colour via its list style."""
        lst = ph.text_frame._txBody.find(qn("a:lstStyle"))
        lvl = lst.find(qn("a:lvl1pPr"))
        if lvl is None:
            lvl = lst.makeelement(qn("a:lvl1pPr"), {}); lst.insert(0, lvl)
        lvl.set("algn", "l")  # brand: left-aligned, flush with the title
        lvl.set("marL", "0"); lvl.set("indent", "0")
        if lvl.find(qn("a:buNone")) is None and lvl.find(qn("a:buChar")) is None:
            lvl.append(parse_xml(f'<a:buNone {nsdecls("a")}/>'))
        d = lvl.find(qn("a:defRPr"))
        if d is None:
            d = lvl.makeelement(qn("a:defRPr"), {}); lvl.append(d)
        for old in d.findall(qn("a:solidFill")):
            d.remove(old)
        d.insert(0, parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{TOK[tok]}"/></a:solidFill>'))
        ph.text_frame.paragraphs[0].alignment = None

    def picture_xml(slide_like, image, left, top, height):
        """Add a picture to a layout/master (python-pptx only adds them to slides)."""
        tmp = prs.slides.add_slide(prs.slide_layouts[6])
        pic = tmp.shapes.add_picture(str(image), left, top, height=height)
        _, rid = slide_like.part.get_or_add_image_part(str(image))
        el = copy.deepcopy(pic._element)
        el.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
        slide_like.shapes._spTree.append(el)
        prs.part.drop_rel(prs.slides._sldIdLst[-1].rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])
        return el

    # Every content slide: colour logo bottom-right, Deep Teal accent rule top-left.
    picture_xml(master, LOGO, W - Inches(1.9), H - Inches(0.75), Inches(0.45))
    tmp = prs.slides.add_slide(prs.slide_layouts[6])  # shapes can't be added to a master directly
    rule = tmp.shapes.add_shape(1, Inches(0.6), Inches(0.42), Inches(0.75), Emu(38100))
    rule.fill.solid(); rule.fill.fore_color.rgb = rgb("--ha-teal-deep"); rule.line.fill.background()
    master.shapes._spTree.append(copy.deepcopy(rule._element))
    prs.part.drop_rel(prs.slides._sldIdLst[-1].rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])

    # Dark layouts (Title Slide, Section Header): Charcoal ground, white logo, no master art.
    for idx in (0, 2):
        lyt = master.slide_layouts[idx]
        lyt._element.set("showMasterSp", "0")
        lyt.background.fill.solid(); lyt.background.fill.fore_color.rgb = rgb("--ha-charcoal")
        picture_xml(lyt, LOGO_WHITE, Inches(0.6), H - Inches(1.1), Inches(0.55))
        for ph, tok in zip(lyt.placeholders, ("--ha-white", "--ha-ash")):
            text_colour(ph, tok)

    def eyebrow(slide, text, dark=False):
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(8), Inches(0.4))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = text.upper()
        r.font.name, r.font.size, r.font.bold = HEAD, Pt(11), True
        r.font.color.rgb = rgb("--ha-ash" if dark else "--ha-teal-deep")
        r.font._rPr.set("spc", "260")  # wide eyebrow tracking

    s = prs.slides.add_slide(master.slide_layouts[0])
    s.shapes.title.text = "Presentation title in Manrope"
    s.placeholders[1].text = "Subtitle or date · Hawaiʻi Appleseed"

    s = prs.slides.add_slide(master.slide_layouts[1])
    eyebrow(s, "Eyebrow label")
    s.shapes.title.text = "Lead with the finding, not the topic"
    body = s.placeholders[1].text_frame
    body.text = "Body text is Poppins in Charcoal. One Deep Teal accent per slide."
    for t in ("Plenty of white space; left-aligned text.", "Spell Hawaiʻi with the ʻokina (U+02BB)."):
        body.add_paragraph().text = t

    s = prs.slides.add_slide(master.slide_layouts[5])
    eyebrow(s, "Chart")
    s.shapes.title.text = "Charts use the brand palette in order"
    cd = CategoryChartData()
    cd.categories = ["2022", "2023", "2024", "2025"]
    for name, vals in (("Series 1", (12, 14, 15, 17)), ("Series 2", (9, 10, 12, 11)), ("Series 3", (6, 8, 7, 9))):
        cd.add_series(name, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.8), Inches(11.5), Inches(4.6), cd).chart
    ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(12), rgb("--ha-charcoal")
    for i, ser in enumerate(ch.series, 1):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(f"--ha-chart-cat-{i}")
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb("--ha-ash")
    ch.value_axis.format.line.fill.background()
    src = s.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(9), Inches(0.4)).text_frame.paragraphs[0]
    r = src.add_run(); r.text = "Source: name the source under every chart."
    r.font.name, r.font.size, r.font.italic, r.font.color.rgb = BODY, Pt(11), True, rgb("--ha-slate")

    s = prs.slides.add_slide(master.slide_layouts[2])
    s.shapes.title.text = "Closing or section slide"
    s.placeholders[1].text = "hawaiiappleseed.org"
    for ph in s.placeholders:
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT

    prs.core_properties.title, prs.core_properties.author = "Hawaiʻi Appleseed slide template", "Hawaiʻi Appleseed"
    prs.save(path)


# ---------------- Word ----------------

def build_docx(path: Path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    rgb = lambda t: RGBColor.from_string(TOK[t])
    doc = Document()
    retheme(doc.part, RT.THEME)
    ft = doc.part.part_related_by(RT.FONT_TABLE)  # tell Word to substitute Arial, not a serif
    ft._blob = ft.blob.replace(b"</w:fonts>", b"".join(
        f'<w:font w:name="{n}"><w:altName w:val="Arial"/><w:charset w:val="00"/><w:family w:val="swiss"/>'
        f'<w:pitch w:val="variable"/></w:font>'.encode() for n in (HEAD, BODY)) + b"</w:fonts>")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)

    def font(style, name, size, colour, bold=None, italic=None):
        f = style.font
        f.name, f.size, f.color.rgb = name, Pt(size), rgb(colour)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rfonts.attrib.pop(qn(attr), None)
        if bold is not None:
            f.bold = bold
        if italic is not None:
            f.italic = italic

    PPR = ("pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc").split()
    RPR = ("rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd").split()

    def insert(parent, el, order):
        """Insert el where the OOXML schema wants it (Word rejects out-of-order children)."""
        name = el.tag.split("}")[1]
        later = set(order[order.index(name) + 1:])
        for i, child in enumerate(parent):
            if child.tag.split("}")[1] in later:
                parent.insert(i, el)
                return el
        parent.append(el)
        return el

    def border(style, side, colour, eighths):
        ppr = style.element.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = insert(ppr, OxmlElement("w:pBdr"), PPR)
        b = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", str(eighths)), ("space", "8"), ("color", TOK[colour])):
            b.set(qn(f"w:{k}"), v)
        bdr.append(b)

    def shade(style, colour):
        shd = OxmlElement("w:shd")
        for k, v in (("val", "clear"), ("color", "auto"), ("fill", TOK[colour])):
            shd.set(qn(f"w:{k}"), v)
        insert(style.element.get_or_add_pPr(), shd, PPR)

    st = doc.styles
    font(st["Normal"], BODY, 10.5, "--ha-charcoal")
    st["Normal"].paragraph_format.line_spacing = 1.3
    st["Normal"].paragraph_format.space_after = Pt(8)
    for name, size, colour, before in (("Title", 28, "--ha-charcoal", 0), ("Heading 1", 18, "--ha-charcoal", 18),
                                       ("Heading 2", 14, "--ha-teal-deep", 14), ("Heading 3", 11.5, "--ha-charcoal", 12)):
        font(st[name], HEAD, size, colour, bold=True)
        st[name].paragraph_format.space_before, st[name].paragraph_format.space_after = Pt(before), Pt(6)
    title_ppr = st["Title"].element.get_or_add_pPr()
    for b in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(b)
    font(st["Subtitle"], BODY, 13, "--ha-teal-deep", italic=False)
    font(st["Quote"], BODY, 13, "--ha-teal-deep", italic=True)
    border(st["Quote"], "left", "--ha-teal", 24)
    st["Quote"].paragraph_format.left_indent = Inches(0.3)

    eb = st.add_style("Eyebrow", WD_STYLE_TYPE.PARAGRAPH)
    eb.base_style, eb.next_paragraph_style = st["Normal"], st["Heading 1"]
    font(eb, HEAD, 8.5, "--ha-teal-deep", bold=True)
    eb.font.all_caps = True
    sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "36"); insert(eb.element.get_or_add_rPr(), sp, RPR)
    eb.paragraph_format.space_after = Pt(2)

    co = st.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    co.base_style = st["Normal"]
    shade(co, "--ha-ash-light")
    border(co, "left", "--ha-teal-deep", 36)
    co.paragraph_format.left_indent = co.paragraph_format.right_indent = Inches(0.15)

    sn = st.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    sn.base_style = st["Normal"]
    font(sn, BODY, 8.5, "--ha-slate", italic=True)

    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(LOGO), height=Inches(0.45))
    fp = sec.footer.paragraphs[0]
    fp.style = sn
    fp.add_run(f"{ORG}  ·  hawaiiappleseed.org  ·  ")
    run = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)

    doc.add_paragraph("Report or brief", style="Eyebrow")
    doc.add_paragraph("Document title in Manrope", style="Title")
    doc.add_paragraph("Subtitle or date · Poppins, Deep Teal", style="Subtitle")
    doc.add_paragraph("Section eyebrow", style="Eyebrow")
    doc.add_heading("Heading 1 names the finding", level=1)
    doc.add_paragraph("Body text is Poppins 10.5 pt in Charcoal. Lead with the number or the finding; "
                      "spell Hawaiʻi with the ʻokina (U+02BB). No serif fonts, no bright colours.")
    doc.add_heading("Heading 2 in Deep Teal", level=2)
    doc.add_paragraph("Callout style: Ash Light fill with a Deep Teal left rule, for the one thing a reader "
                      "should take away from the section.", style="Callout")
    doc.add_paragraph("Pull-quote style: Poppins italic in Deep Teal with a Teal rule.", style="Quote")
    doc.add_heading("Heading 3 for sub-points", level=3)
    doc.add_paragraph("Source: the Source Note style goes under every chart and table.", style="Source Note")
    doc.core_properties.title, doc.core_properties.author = "Hawaiʻi Appleseed document template", "Hawaiʻi Appleseed"
    doc.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_pptx(OUT / "appleseed-slides.pptx")
    build_docx(OUT / "appleseed-document.docx")
    print(f"wrote {OUT.relative_to(REPO)}/appleseed-slides.pptx, appleseed-document.docx")


if __name__ == "__main__":
    main()
